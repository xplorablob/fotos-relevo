"""
Aplicador de marca de agua para fotos de relevo (Xplora).

Lee un Excel con columnas FOTO, FECHA, HORA, RAZON SOCIAL, DISTANCIA (y DNI opcional),
descarga cada foto, le pone la franja gris transparente inferior con letras blancas
(igual que la marca de agua del Xplora App) y guarda la foto en la carpeta de salida.
Devuelve el mismo Excel con la columna FOTO_MARCA_AGUA (link público nuevo).

Uso:
  python marca_agua.py ENTRADA.xlsx --base-url https://USUARIO.github.io/fotos-relevo/fotos/ \
         --out-dir docs/fotos --out-excel salida/ENTRADA_CON_LINKS.xlsx
"""
import argparse, io, os, sys, datetime as dt
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urlparse

import requests
import openpyxl
from PIL import Image, ImageDraw, ImageFont, ImageOps

AQUI = os.path.dirname(os.path.abspath(__file__))
FUENTES = [
    os.path.join(AQUI, "fuentes", "Roboto-Regular.ttf"),
    "/usr/share/fonts/truetype/roboto/unhinted/RobotoTTF/Roboto-Regular.ttf",
    "/usr/share/fonts/truetype/roboto/hinted/Roboto-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
]
SUFIJO = "MA"          # se agrega al final del nombre: ...8406.jpg -> ...8406MA.jpg
ANCHO_MAX = 1200       # ancho final de la foto (se ve nítida y pesa ~150 KB)
CALIDAD_JPG = 80

# Proporciones medidas de las fotos de referencia (relativas al ancho de la foto)
TAM_LETRA = 0.0450     # alto de letra
INTERLINEA = 0.0510    # distancia entre renglones
MARGEN_X = 0.008       # margen izquierdo del texto
PAD_SUP = 0.020        # aire arriba del primer renglón
PAD_DER = 0.008        # aire a la derecha del renglón más largo
GRIS = (150, 150, 150, 115)  # gris transparente


def fuente(px):
    for f in FUENTES:
        if os.path.exists(f):
            return ImageFont.truetype(f, px)
    return ImageFont.load_default(px)


def fmt_fecha(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    for f in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s[:10], f).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


def fmt_hora(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M:%S")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, (int, float)):  # fracción de día de Excel
        seg = round(float(v) % 1 * 86400)
        return f"{seg//3600:02d}:{seg%3600//60:02d}:{seg%60:02d}"
    return str(v or "").strip()


def fmt_distancia(v):
    s = str(v if v is not None else "").strip()
    if not s:
        return ""
    if s.replace(".", "", 1).isdigit():
        s += " mts"
    return s if s.endswith(".") else s + "."


def lineas_de(fila):
    return [
        f"Dia: {fmt_fecha(fila.get('FECHA'))} Hora: {fmt_hora(fila.get('HORA'))}",
        f"DNI: {str(fila.get('DNI') or '').strip()}",
        str(fila.get("RAZON SOCIAL") or "").strip(),
        f"Distancia: {fmt_distancia(fila.get('DISTANCIA'))}",
    ]


def poner_marca(img, lineas):
    img = ImageOps.exif_transpose(img).convert("RGB")
    if img.width > ANCHO_MAX:
        img = img.resize((ANCHO_MAX, round(img.height * ANCHO_MAX / img.width)), Image.LANCZOS)
    W, H = img.size
    escala = 1.0
    f = fuente(round(W * TAM_LETRA))
    disponible = W * (1 - MARGEN_X - PAD_DER)
    largo = max(f.getlength(t) for t in lineas)
    if largo > disponible:  # razón social muy larga: se achica la letra para que entre
        escala = disponible / largo
        f = fuente(max(8, round(W * TAM_LETRA * escala)))
    inter = W * INTERLINEA * escala
    asc, desc = f.getmetrics()
    alto_franja = round(W * PAD_SUP * escala + inter * (len(lineas) - 1) + asc + desc * 0.6)
    ancho_txt = max(f.getlength(t) for t in lineas)
    x0 = round(W * MARGEN_X)
    ancho_franja = min(W, round(x0 + ancho_txt + W * PAD_DER))
    y0 = H - alto_franja

    capa = Image.new("RGBA", img.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(capa)
    d.rectangle([0, y0, ancho_franja, H], fill=GRIS)
    for i, t in enumerate(lineas):
        d.text((x0, y0 + W * PAD_SUP * escala - asc * 0.22 + i * inter), t, font=f, fill=(255, 255, 255, 255))
    return Image.alpha_composite(img.convert("RGBA"), capa).convert("RGB")


def nombre_nuevo(url):
    base = os.path.basename(urlparse(url).path)
    raiz, ext = os.path.splitext(base)
    return f"{raiz}{SUFIJO}{ext or '.jpg'}"


def procesar(fila, out_dir, sesion):
    url = str(fila.get("FOTO") or "").strip()
    if not url.startswith("http"):
        return None, "sin URL"
    nombre = nombre_nuevo(url)
    destino = os.path.join(out_dir, nombre)
    if os.path.exists(destino):
        return nombre, "ya existía"
    r = sesion.get(url, timeout=60)
    r.raise_for_status()
    img = poner_marca(Image.open(io.BytesIO(r.content)), lineas_de(fila))
    img.save(destino, "JPEG", quality=CALIDAD_JPG, optimize=True)
    return nombre, "ok"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--base-url", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--out-excel", required=True)
    ap.add_argument("--hilos", type=int, default=16)
    a = ap.parse_args()
    base = a.base_url.rstrip("/") + "/"
    os.makedirs(a.out_dir, exist_ok=True)
    os.makedirs(os.path.dirname(a.out_excel) or ".", exist_ok=True)

    wb = openpyxl.load_workbook(a.excel)  # conserva fórmulas y formato
    ws = wb.worksheets[0]
    enc = [str(c.value or "").strip().upper() for c in ws[1]]
    if "FOTO" not in enc:
        sys.exit("No encontré la columna FOTO en la primera hoja")
    col_nueva = enc.index("FOTO_MARCA_AGUA") + 1 if "FOTO_MARCA_AGUA" in enc else len(enc) + 1
    col_estado = enc.index("ESTADO_MARCA") + 1 if "ESTADO_MARCA" in enc else col_nueva + 1
    ws.cell(1, col_nueva, "FOTO_MARCA_AGUA")
    ws.cell(1, col_estado, "ESTADO_MARCA")

    vals = openpyxl.load_workbook(a.excel, data_only=True).worksheets[0]
    filas = []
    for r in range(2, ws.max_row + 1):
        fila = {enc[c]: vals.cell(r, c + 1).value for c in range(len(enc))}
        if any(v not in (None, "") for v in fila.values()):
            filas.append((r, fila))

    sesion = requests.Session()
    ok = err = 0
    with ThreadPoolExecutor(a.hilos) as ex:
        futs = {ex.submit(procesar, f, a.out_dir, sesion): r for r, f in filas}
        for i, fu in enumerate(as_completed(futs), 1):
            r = futs[fu]
            try:
                nombre, estado = fu.result()
                ws.cell(r, col_nueva, base + nombre if nombre else None)
                ws.cell(r, col_estado, estado)
                ok += nombre is not None
            except Exception as e:
                ws.cell(r, col_estado, f"error: {e}"[:200])
                err += 1
            if i % 200 == 0:
                print(f"{i}/{len(filas)} procesadas", flush=True)
    wb.save(a.out_excel)
    print(f"Listo: {ok} fotos con marca, {err} errores. Excel: {a.out_excel}")


if __name__ == "__main__":
    main()
